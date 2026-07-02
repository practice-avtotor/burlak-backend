from __future__ import annotations

import io
import json
import logging
import os
import traceback
import zipfile
from typing import Any

import openpyxl

from app.core.config import get_settings
from app.db import sync_repository
from app.services.card_parser_service import (
    CardParserService,
    MLCardParseResult,
    ParsedPart,
)
from app.services.heuristic_analyzer import extract_card_number_from_filepath
from app.services.splitter import CardSplitter
from app.services.structure_adapter import StructureAdapter
from app.services.xls_converter import convert_xls_to_xlsx

logger = logging.getLogger(__name__)


def extract_card_number(card_path: str) -> str:
    """Extract card number from card path/filename.

    Fallback to basename without extension.
    """
    return extract_card_number_from_filepath(card_path)


class CardProcessingService:
    """Service to process a single operational card including splitting, translation, and Excel output generation."""

    def __init__(
        self,
        job_id: int,
        ml_client: StructureAdapter | None = None,
    ) -> None:
        self.job_id = job_id
        self.settings = get_settings()
        self._ml_client = ml_client

    def process_card(self, card_path: str) -> None:
        logger.info(f"Processing card for job {self.job_id}: {card_path}")

        # Define directories
        storage_root = self.settings.storage_path
        job_dir = os.path.join(storage_root, str(self.job_id))
        translated_cards_dir = os.path.join(job_dir, "translated_cards")
        card_materials_dir = os.path.join(job_dir, "card_materials")

        # 1. Retrieve job details & files
        _, archive_path = sync_repository.get_job_files(self.job_id)
        if not archive_path or not os.path.exists(archive_path):
            raise FileNotFoundError(f"Archive file not found: {archive_path}")

        mapping_config = sync_repository.get_mapping_config(self.job_id)
        if not mapping_config:
            raise ValueError(f"Mapping configuration not found for job {self.job_id}")

        # 2. Open archive and stream card data
        with zipfile.ZipFile(archive_path, "r") as zf:
            try:
                card_bytes = zf.read(card_path)
            except KeyError:
                raise FileNotFoundError(f"Card {card_path} not found in zip archive")

        filename = os.path.basename(card_path)

        # 2b. XLS conversion to XLSX via LibreOffice if needed
        if card_path.lower().endswith(".xls"):
            logger.info("Converting legacy XLS file to XLSX: %s", card_path)
            temp_xls_path = os.path.join(
                job_dir, f"temp_{self.job_id}_{os.path.basename(card_path)}"
            )
            os.makedirs(job_dir, exist_ok=True)
            with open(temp_xls_path, "wb") as f:
                f.write(card_bytes)

            converted_xlsx_path = convert_xls_to_xlsx(temp_xls_path, output_dir=job_dir)
            if not converted_xlsx_path or not os.path.exists(converted_xlsx_path):
                raise ValueError(f"Failed to convert XLS file to XLSX: {card_path}")

            with open(converted_xlsx_path, "rb") as f:
                card_bytes = f.read()

            # Clean up temp converted files
            try:
                os.remove(temp_xls_path)
                os.remove(converted_xlsx_path)
            except OSError:
                pass

            filename = os.path.splitext(filename)[0] + ".xlsx"

        # 3. Classify card
        parser = CardParserService(mapping_config)
        classification = parser.classify(filename)
        logger.info(f"Card {card_path} classified as '{classification}'")

        # Destination paths
        dest_xlsx_path = os.path.join(
            translated_cards_dir, os.path.dirname(card_path), filename
        )
        os.makedirs(os.path.dirname(dest_xlsx_path), exist_ok=True)

        card_materials_path = os.path.join(
            card_materials_dir, os.path.dirname(card_path), f"{filename}.json"
        )
        os.makedirs(os.path.dirname(card_materials_path), exist_ok=True)

        if classification != "operational_card":
            # Non-operational card (service or unknown). Save original directly and skip translation/parsing.
            with open(dest_xlsx_path, "wb") as f:
                f.write(card_bytes)

            # Create an empty materials file so aggregate task doesn't fail
            with open(card_materials_path, "w", encoding="utf-8") as f:
                json.dump([], f)
            return

        # 3b. Handle Sheet Splitting for multi-sheet cards
        temp_source_path = os.path.join(job_dir, f"split_src_{self.job_id}_{filename}")
        os.makedirs(job_dir, exist_ok=True)
        with open(temp_source_path, "wb") as f:
            f.write(card_bytes)

        try:
            wb_check = openpyxl.load_workbook(temp_source_path, read_only=True)
            sheet_names = wb_check.sheetnames
            wb_check.close()
        except Exception as e:
            logger.warning(
                "Failed to check sheets with openpyxl, treating as single sheet: %s", e
            )
            sheet_names = []

        # Find sheets to split (exclude service sheets)
        sheets_to_split = []
        for sn in sheet_names:
            sn_lower = sn.lower()
            is_svc_sheet = any(
                kw.lower() in sn_lower for kw in parser.service_keywords
            )
            if not is_svc_sheet:
                sheets_to_split.append(sn)

        split_files = []
        if len(sheets_to_split) > 1:
            logger.info(
                "Multi-sheet card detected. Splitting %d sheets: %s",
                len(sheets_to_split),
                sheets_to_split,
            )
            split_dir = os.path.join(
                job_dir, "split_cards", os.path.splitext(filename)[0]
            )
            splitter = CardSplitter(max_workers=1)
            try:
                split_files = splitter.split_file(
                    temp_source_path,
                    split_dir,
                    sheets_to_split,
                    file_label=os.path.splitext(filename)[0],
                )
            except Exception as e:
                logger.error("Failed to split card %s: %s", filename, e, exc_info=True)

        # Cleanup temp source file
        try:
            os.remove(temp_source_path)
        except OSError:
            pass

        # 4. Parse the operational card (single sheet or all split parts)
        if split_files:
            # Process each split file as a separate sub-card
            for sf in split_files:
                sf_name = os.path.basename(sf)
                logger.info("Processing split operational card: %s", sf_name)

                with open(sf, "rb") as f:
                    sf_bytes = f.read()

                sf_parse_result = parser.parse_card(sf_bytes, sf_name)
                if sf_parse_result.error:
                    logger.warning("Split file parse error: %s", sf_parse_result.error)
                    continue

                sf_card_no = extract_card_number(sf_name)
                sf_dest_xlsx_path = os.path.join(
                    translated_cards_dir, os.path.dirname(card_path), sf_name
                )
                sf_card_materials_path = os.path.join(
                    card_materials_dir, os.path.dirname(card_path), f"{sf_name}.json"
                )
                self._translate_and_save_card(
                    parts=sf_parse_result.parts,
                    dest_path=sf_dest_xlsx_path,
                    materials_path=sf_card_materials_path,
                    card_bytes=sf_bytes,
                    card_no=sf_card_no,
                    name_col=parser.name_col,
                )

                # Cleanup temp split file
                try:
                    os.remove(sf)
                except OSError:
                    pass
        else:
            # Single-sheet mode
            parse_result = parser.parse_card(card_bytes, filename)
            if parse_result.error:
                raise ValueError(parse_result.error)

            card_no = extract_card_number(card_path)

            # Check if this is a multi-card result (multiple cards in one sheet)
            card_boundaries = parse_result.card_boundaries
            if card_boundaries and len(card_boundaries) > 1:
                # Multi-card: split parts by card boundaries and process each sub-card
                logger.info(
                    "Multi-card sheet detected: %d cards in %s",
                    len(card_boundaries),
                    filename,
                )
                self._process_multi_card_parts(
                    job_id=self.job_id,
                    card_path=card_path,
                    filename=filename,
                    card_bytes=card_bytes,
                    parse_result=parse_result,
                    card_boundaries=card_boundaries,
                    card_no=card_no,
                    mapping_config=mapping_config,
                    translated_cards_dir=translated_cards_dir,
                    card_materials_dir=card_materials_dir,
                    name_col=parser.name_col,
                )
            else:
                # Single card — original logic
                self._translate_and_save_card(
                    parts=parse_result.parts,
                    dest_path=dest_xlsx_path,
                    materials_path=card_materials_path,
                    card_bytes=card_bytes,
                    card_no=card_no,
                    name_col=parser.name_col,
                )

    def _translate_and_save_card(
        self,
        parts: list[ParsedPart],
        dest_path: str,
        materials_path: str,
        card_bytes: bytes,
        card_no: str,
        name_col: int,
    ) -> None:
        """Translate Chinese names, save translated workbook, and write materials JSON.

        Shared helper used by split-file, single-card, and multi-card paths.
        """
        unique_texts = {p.name for p in parts if p.name}
        translations: dict[str, str] = {}
        if unique_texts:
            try:
                if self._ml_client:
                    translations = self._ml_client.translate_batch(
                        list(unique_texts), target_lang="ru"
                    )
                else:
                    with StructureAdapter(self.settings.ml_service_url) as ml_client:
                        translations = ml_client.translate_batch(
                            list(unique_texts), target_lang="ru"
                        )
            except Exception as e:
                logger.warning(f"Translation batch failed: {e}. Using original text.")

        # Build parts list for materials JSON
        json_parts = [
            {
                "card_no": card_no,
                "sheet_name": p.source_sheet,
                "row_index": p.row,
                "part_no": p.part_number,
                "name_cn": p.name,
                "name_ru": translations.get(p.name, p.name),
                "qty": p.quantity,
            }
            for p in parts
        ]

        # Save translated workbook
        os.makedirs(os.path.dirname(dest_path), exist_ok=True)
        wb = openpyxl.load_workbook(io.BytesIO(card_bytes), data_only=False)
        try:
            for p in parts:
                if p.name and name_col > 0:
                    ws = wb[p.source_sheet]
                    cell = ws.cell(row=p.row, column=name_col)
                    if cell.font and cell.font.strike:
                        continue
                    cell.value = translations.get(p.name, p.name)
            wb.save(dest_path)
        finally:
            wb.close()

        # Save parts list as JSON
        os.makedirs(os.path.dirname(materials_path), exist_ok=True)
        with open(materials_path, "w", encoding="utf-8") as f:
            json.dump(json_parts, f, ensure_ascii=False, indent=2)

    def write_error_file(self, card_path: str, exc: Exception) -> None:
        """Write error materials file containing exception detail."""
        storage_root = self.settings.storage_path
        job_dir = os.path.join(storage_root, str(self.job_id))
        card_materials_dir = os.path.join(job_dir, "card_materials")

        try:
            error_materials_path = os.path.join(
                card_materials_dir, f"{card_path}_error.json"
            )
            os.makedirs(os.path.dirname(error_materials_path), exist_ok=True)
            with open(error_materials_path, "w", encoding="utf-8") as f:
                json.dump(
                    {
                        "card_path": card_path,
                        "error": str(exc),
                        "traceback": traceback.format_exc(),
                    },
                    f,
                    ensure_ascii=False,
                    indent=2,
                )
        except Exception as write_err:
            logger.error(f"Failed to write error marker for {card_path}: {write_err}")

    def _process_multi_card_parts(
        self,
        job_id: int,
        card_path: str,
        filename: str,
        card_bytes: bytes,
        parse_result: MLCardParseResult,
        card_boundaries: list[tuple[str, int, int]],
        card_no: str,
        mapping_config: dict[str, Any],
        translated_cards_dir: str,
        card_materials_dir: str,
        name_col: int = 0,
    ) -> None:
        """Process a multi-card sheet by splitting parts into sub-cards.

        Each sub-card gets its own translated XLSX and materials JSON.
        """
        # Group parts by which card boundary they fall into.
        # Boundaries now include sheet_name to prevent cross-sheet mixing.
        sub_card_parts: list[list[ParsedPart]] = [[] for _ in card_boundaries]

        for p in parse_result.parts:
            assigned = False
            for idx, (sheet_name, start_row, end_row) in enumerate(card_boundaries):
                if p.source_sheet == sheet_name and start_row <= p.row <= end_row:
                    sub_card_parts[idx].append(p)
                    assigned = True
                    break
            if not assigned and sub_card_parts:
                sub_card_parts[-1].append(p)
        for idx, parts_group in enumerate(sub_card_parts):
            if not parts_group:
                logger.warning("No parts found for sub-card %d, skipping", idx)
                continue

            # Generate sub-card identifier
            sub_suffix = f"_part{idx + 1}"
            sub_name = f"{os.path.splitext(filename)[0]}{sub_suffix}.xlsx"
            sub_card_no = f"{card_no}{sub_suffix}"

            logger.info(
                "Processing sub-card %d/%d: %s (%d parts)",
                idx + 1,
                len(card_boundaries),
                sub_name,
                len(parts_group),
            )

            sub_dest_xlsx_path = os.path.join(
                translated_cards_dir, os.path.dirname(card_path), sub_name
            )
            sub_card_materials_path = os.path.join(
                card_materials_dir,
                os.path.dirname(card_path),
                f"{sub_name}.json",
            )
            self._translate_and_save_card(
                parts=parts_group,
                dest_path=sub_dest_xlsx_path,
                materials_path=sub_card_materials_path,
                card_bytes=card_bytes,
                card_no=sub_card_no,
                name_col=name_col,
            )
