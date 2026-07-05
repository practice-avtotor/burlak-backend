from typing import Any

from openpyxl import Workbook

from app.services.heuristic_analyzer import (
    HeuristicAnalyzer,
    clean_cell_text,
    extract_card_number_from_filepath,
)


def _make_ws(data: list[list[Any | None]]) -> Any:
    wb = Workbook()
    ws = wb.active
    for r_idx, row in enumerate(data, 1):
        for c_idx, val in enumerate(row, 1):
            if val is not None:
                ws.cell(row=r_idx, column=c_idx, value=val)
    return ws


def _make_ws_from_dict(rows: dict[int, dict[int, Any]]) -> Any:
    wb = Workbook()
    ws = wb.active
    for r, cols in rows.items():
        for c, val in cols.items():
            ws.cell(row=r, column=c, value=val)
    return ws


class TestCleanCellText:
    def test_clean_normal(self):
        assert clean_cell_text("Normal Text") == "Normal Text"
        assert clean_cell_text(None) == ""

    def test_clean_xml_hex(self):
        assert clean_cell_text("Part_x0020_No") == "Part"

    def test_clean_newline(self):
        assert clean_cell_text("Part\nName") == "Part"
        assert clean_cell_text("Part\r\nName") == "Part"


class TestExtractCardNumberFromFilepath:
    def test_filename_patterns(self):
        assert (
            extract_card_number_from_filepath("/path/to/ABC1L-17-AS-04001.xlsx")
            == "ABC1L-17-AS-04001"
        )
        assert (
            extract_card_number_from_filepath("/path/to/A001_part.xlsx") == "A001_part"
        )
        assert (
            extract_card_number_from_filepath("/path/to/123-AS-456.xlsx")
            == "123-AS-456"
        )
        assert extract_card_number_from_filepath("/path/to/other.xlsx") == "other"


class TestGetCellValue:
    def test_openpyxl_worksheet(self):
        ws = _make_ws([["Hello", "World"]])
        assert HeuristicAnalyzer.get_cell_value(ws, 1, 1) == "Hello"
        assert HeuristicAnalyzer.get_cell_value(ws, 1, 2) == "World"
        assert HeuristicAnalyzer.get_cell_value(ws, 2, 1) is None

    def test_custom_excel_sheet(self):
        class MockExcelSheet:
            def cell_value(self, row, col):
                return f"R{row}C{col}"

        ms = MockExcelSheet()
        assert HeuristicAnalyzer.get_cell_value(ms, 3, 5) == "R3C5"

    def test_exception_safety(self):
        class BrokenSheet:
            @property
            def cell(self):
                raise RuntimeError("Broken")

        bs = BrokenSheet()
        assert HeuristicAnalyzer.get_cell_value(bs, 1, 1) is None


class TestIsCellStrike:
    def test_not_strike(self):
        ws = _make_ws([["Text"]])
        assert not HeuristicAnalyzer.is_cell_strike(ws, 1, 1)

    def test_strike(self):
        ws = _make_ws([["Text"]])
        ws.cell(row=1, column=1).font = ws.cell(row=1, column=1).font.copy(strike=True)
        assert HeuristicAnalyzer.is_cell_strike(ws, 1, 1)


class TestGetStrikeRows:
    def test_strike_rows(self):
        ws = _make_ws([["Part1"], ["Part2"], ["Part3"]])
        # Set row 2 cell to strike
        ws.cell(row=2, column=1).font = ws.cell(row=2, column=1).font.copy(strike=True)
        assert HeuristicAnalyzer.get_strike_rows(ws, range(1, 4), [1]) == {2}


class TestFindPartTable:
    def test_swm_card_part_table(self):
        rows: dict[int, dict[int, str]] = {
            5: {17: "序号", 18: "零部件代号"},
            6: {18: "P001"},
        }
        ws = _make_ws_from_dict(rows)
        result = HeuristicAnalyzer.find_part_table(ws)
        assert result is not None
        hr, pn, qty, name = result
        assert hr == 5
        assert pn == 18

    def test_t1l_card_part_table(self):
        rows: dict[int, dict[int, str]] = {
            2: {1: "物料编码", 2: "零件名称", 3: "数量"},
            3: {1: "P001", 2: "Part1", 3: "1"},
        }
        ws = _make_ws_from_dict(rows)
        result = HeuristicAnalyzer.find_part_table(ws)
        assert result is not None
        hr, pn, qty, name = result
        assert hr == 2
        assert pn == 1
        assert name == 2
        assert qty == 3


class TestExtractOperationName:
    def test_finds_operation_name(self):
        rows: dict[int, dict[int, str]] = {
            2: {4: "安装左前门线束"},
            5: {17: "序号", 18: "零部件代号"},
        }
        ws = _make_ws_from_dict(rows)
        name = HeuristicAnalyzer.extract_operation_name(ws, 5)
        assert name == "安装左前门线束"
