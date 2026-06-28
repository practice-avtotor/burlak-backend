import sqlite3
import zipfile
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]

from app.worker.celery_app import celery_app
from app.worker.tasks.unpack import unpack


def test_celery_pipeline_success(temp_db_path: str, mock_storage_path: Path) -> None:
    """Verifies the complete Celery task pipeline end-to-end in eager mode."""
    # 1. Enable eager execution mode for Celery
    celery_app.conf.task_always_eager = True
    celery_app.conf.task_eager_propagates = True

    # 2. Setup job files and directories
    job_id = 1
    job_dir = mock_storage_path / str(job_id)
    job_dir.mkdir(parents=True, exist_ok=True)

    bom_path = job_dir / "bom.xlsx"
    archive_path = job_dir / "archive.zip"

    # Create dummy BOM
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "总装BOM"
    ws.append(["序号", "零件号", "零件名称", "用量", "备注"])
    ws.append([1, "PART-001", "螺栓M6×20", 10, "BOM item"])
    wb.save(bom_path)

    # Create dummy ZIP containing an operational card and a service card
    with zipfile.ZipFile(archive_path, "w") as zf:
        # Create an operational card
        card_wb = openpyxl.Workbook()
        card_ws = card_wb.active
        card_ws.title = "Sheet1"
        card_ws.append(["序号", "零件号", "名称", "数量", "备注"])
        card_ws.append([1, "PART-001", "螺栓M6×20", 10, "Card item"])

        import io

        card_bio = io.BytesIO()
        card_wb.save(card_bio)
        zf.writestr("SQRT1L-17-AS-04001.xlsx", card_bio.getvalue())

        # Create a service file (using "Cover" keyword to trigger service classification)
        service_wb = openpyxl.Workbook()
        service_ws = service_wb.active
        service_ws.title = "Cover Page"
        service_ws.append(["Document Title"])
        service_ws.append(["Manufacturing Cover Info"])

        service_bio = io.BytesIO()
        service_wb.save(service_bio)
        zf.writestr("Assembly_Cover.xlsx", service_bio.getvalue())

    # 3. Create job record in the DB
    from datetime import UTC, datetime

    now = datetime.now(UTC).isoformat()
    conn = sqlite3.connect(temp_db_path)
    conn.execute(
        """
        INSERT INTO jobs (id, status, stage, bom_path, archive_path, processed, failed, total, bom_uploaded, archive_uploaded, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            job_id,
            "processing",
            "unpacking",
            str(bom_path),
            str(archive_path),
            0,
            0,
            0,
            True,
            True,
            now,
            now,
        ),
    )
    conn.commit()
    conn.close()

    # 4. Trigger Celery pipeline synchronously with mocked ML service
    from unittest.mock import patch

    mock_mapping_config = {
        "bom": {
            "columns": {"part_no": 2, "qty": 4, "name": 3},
            "table_boundaries": {"header_row": 1, "data_start_row": 2, "end_markers": ["END"]},
        },
        "cards": {
            "columns": {"part_no": 2, "qty": 4, "name": 3},
            "table_boundaries": {"header_row": 1, "data_start_row": 2, "end_markers": ["END"]},
            "file_classification_rules": {
                "operational_card_patterns": [
                    {"type": "filename_regex", "pattern": ".*-AS-.*", "format_group": "card_format_A"}
                ],
                "service_file_patterns": [
                    {"type": "filename_keyword", "keywords": ["Cover", "封面", "目录"]}
                ]
            }
        }
    }

    with patch("app.services.structure_adapter.StructureAdapter.analyze_structure", return_value=mock_mapping_config) as mock_analyze, \
         patch("app.services.structure_adapter.StructureAdapter.translate_batch", return_value={"螺栓M6×20": "Bolt M6x20"}) as mock_translate:
        unpack.delay(job_id)

        # Assert ML client was invoked
        mock_analyze.assert_called_once()

    # 5. Verify database state
    conn = sqlite3.connect(temp_db_path)
    conn.row_factory = sqlite3.Row
    cursor = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,))
    job = cursor.fetchone()

    # Assert job is done
    assert job["status"] == "done"
    assert job["stage"] == "completed"
    assert job["total"] == 2
    assert job["processed"] == 2
    assert job["failed"] == 0

    # Assert card records were created
    cursor = conn.execute(
        "SELECT * FROM cards WHERE job_id = ? ORDER BY card_path", (job_id,)
    )
    cards = cursor.fetchall()
    assert len(cards) == 2
    assert cards[0]["card_path"] == "Assembly_Cover.xlsx"
    assert cards[0]["status"] == "success"
    assert cards[1]["card_path"] == "SQRT1L-17-AS-04001.xlsx"
    assert cards[1]["status"] == "success"
    conn.close()

    # 6. Verify result files exist
    diff_file = job_dir / "diff.xlsx"
    zip_file = job_dir / "translated_cards.zip"

    assert diff_file.exists()
    assert zip_file.exists()

    # Assert temporary folder was cleaned up
    temp_translated_dir = job_dir / "translated_cards"
    assert not temp_translated_dir.exists()
